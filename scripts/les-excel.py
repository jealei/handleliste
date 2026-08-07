#!/usr/bin/env python3
"""Leser arket «Ny» og skriver en ordrett avskrift til src/data/rader.ts.

Skriptet gjør bevisst SÅ LITE som mulig: ingen id-generering, ingen gruppering,
ingen opprydding. Alt slikt hører hjemme i TypeScript, der det kan testes.
Skriptet er en ren avskrift, slik at diffen mot forrige kjøring er lett å lese.

Kjør:  npm run importer
"""
from pathlib import Path
import sys

try:
    import openpyxl
except ImportError:
    sys.exit('Mangler openpyxl. Installer med:  python3 -m pip install openpyxl')

KILDE = Path.home() / 'Downloads' / 'Handleliste Tønsberg.xlsx'
ARK = 'Ny'
FORSTE_DATARAD = 7
MAL = Path(__file__).resolve().parent.parent / 'src' / 'data' / 'rader.ts'

FELT = ['kategori', 'vare', 'merke', 'borHa', 'forpakning', 'kommentar']


def tekst(celle):
    """Gjør en celle om til trimmet tekst. Heltall skal ikke bli «5.0»."""
    verdi = celle.value
    if verdi is None:
        return ''
    if isinstance(verdi, float) and verdi.is_integer():
        return str(int(verdi))
    return str(verdi).strip()


def les_rader(ark):
    rader, kategori = [], None
    for rad in ark.iter_rows(min_row=FORSTE_DATARAD, max_row=ark.max_row, max_col=7):
        a, b, c, d, _kjop, f, g = (tekst(celle) for celle in rad)
        if a:
            kategori = a
        elif b:
            if kategori is None:
                sys.exit(f'Vare uten kategori på rad {rad[0].row}: {b!r}')
            rader.append({
                'kategori': kategori, 'vare': b, 'merke': c,
                'borHa': d, 'forpakning': f, 'kommentar': g,
            })
    return rader


def sitat(verdi):
    """Escaper til et TypeScript-strengliteral med enkle apostrofer."""
    return "'" + verdi.replace('\\', '\\\\').replace("'", "\\'").replace('\n', ' ') + "'"


def skriv(rader):
    linjer = [
        '// GENERERT FIL — IKKE REDIGER FOR HÅND.',
        '// Kjør `npm run importer` for å regenerere fra regnearket.',
        f'// Kilde: {KILDE.name}, arket «{ARK}».',
        '',
        'export type Rad = {',
        '  kategori: string',
        '  vare: string',
        '  merke?: string',
        '  borHa?: string',
        '  forpakning?: string',
        '  kommentar?: string',
        '}',
        '',
        'export const RADER: readonly Rad[] = [',
    ]
    for rad in rader:
        par = [f'{felt}: {sitat(rad[felt])}' for felt in FELT if rad[felt]]
        linjer.append('  { ' + ', '.join(par) + ' },')
    linjer += [']', '']
    MAL.parent.mkdir(parents=True, exist_ok=True)
    MAL.write_text('\n'.join(linjer), encoding='utf-8')


def main():
    if not KILDE.exists():
        sys.exit(f'Finner ikke regnearket: {KILDE}')
    bok = openpyxl.load_workbook(KILDE, data_only=True)
    if ARK not in bok.sheetnames:
        sys.exit(f'Fant ikke arket «{ARK}». Ark i boka: {bok.sheetnames}')
    rader = les_rader(bok[ARK])
    kategorier = list(dict.fromkeys(r['kategori'] for r in rader))
    if not rader:
        sys.exit('Fant ingen varer — sjekk at arkstrukturen ikke er endret.')
    skriv(rader)
    print(f'Skrev {len(rader)} varer i {len(kategorier)} kategorier til {MAL}')
    for navn in kategorier:
        print(f"  {sum(1 for r in rader if r['kategori'] == navn):3d}  {navn}")


if __name__ == '__main__':
    main()
